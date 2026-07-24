"""Finalize và report snapshot bất biến cho kỳ thi."""

import hashlib
import json
from pathlib import Path

from training import constants, time_policy
from training.db import read_connection, write_connection
from training.errors import ErrorCode, TrainingError
from repositories.training_repository import gen_id, write_audit
from services import training_attempt_service as attempts
from services import training_exam_service as exams

REPORT_SCHEMA_VERSION = "1.0"
REPORT_ALGORITHM_VERSION = "v1"


def _checksum(payload):
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _build_payload(db_path, exam_id):
    conn = read_connection(db_path)
    try:
        assignments = [dict(row) for row in conn.execute(
            "SELECT * FROM exam_assignments WHERE exam_event_id=? ORDER BY username", (exam_id,)
        ).fetchall()]
        individual = []
        for assignment in assignments:
            result = conn.execute(
                """SELECT r.raw_score, r.maximum_score, r.percent, r.passed, r.topic_breakdown_json
                   FROM exam_results r
                   JOIN exam_attempts a ON a.id=r.attempt_id
                   WHERE a.assignment_id=?""",
                (assignment["id"],),
            ).fetchone()
            individual.append({
                "assignment": assignment,
                "result": ({**dict(result), "topic_breakdown": json.loads(result["topic_breakdown_json"])}
                           if result else None),
            })
        completed = [row for row in individual if row["result"]]
        return {
            "schema_version": REPORT_SCHEMA_VERSION,
            "exam_id": exam_id,
            "summary": {
                "assigned": len(assignments),
                "completed": len(completed),
                "expired": sum(a["status"] == "expired" for a in assignments),
                "passed": sum(bool(row["result"]["passed"]) for row in completed),
                "failed": sum(not bool(row["result"]["passed"]) for row in completed),
            },
            "individual": individual,
        }
    finally:
        conn.close()


def finalize_exam(db_path, *, unit_code, actor, exam_id):
    """Chốt kỳ thi idempotently; không giữ transaction qua scoring/report build."""
    exam = exams.get_exam(db_path, exam_id)
    if not exam:
        raise TrainingError(ErrorCode.NOT_FOUND, "Kỳ thi không tồn tại", status=404)
    if exam["status"] == constants.ExamStatus.OPEN:
        if time_policy.utc_now_ms() < exam["end_at_ms"]:
            raise TrainingError(
                ErrorCode.EXAM_NOT_OPEN,
                "Kỳ thi vẫn đang trong thời gian làm bài", status=409,
            )
        exams.close_exam(db_path, unit_code=unit_code, actor=actor, exam_id=exam_id)
    elif exam["status"] != constants.ExamStatus.CLOSED:
        raise TrainingError(
            ErrorCode.CONFLICT,
            f"Không thể finalize: trạng thái {exam['status']}", status=409,
        )

    conn = read_connection(db_path)
    try:
        active_attempt_ids = [row["id"] for row in conn.execute(
            """SELECT a.id FROM exam_attempts a
               JOIN exam_assignments x ON x.id=a.assignment_id
               WHERE x.exam_event_id=? AND a.status='active'""",
            (exam_id,),
        ).fetchall()]
    finally:
        conn.close()

    for attempt_id in active_attempt_ids:
        attempts.administratively_submit_attempt(
            db_path, unit_code=unit_code, actor=actor, attempt_id=attempt_id,
            ended_reason="exam_closed",
        )

    conn = write_connection(db_path)
    try:
        conn.execute("BEGIN IMMEDIATE")
        conn.execute(
            """UPDATE exam_assignments SET status='expired'
               WHERE exam_event_id=? AND status='assigned'
                 AND id NOT IN (SELECT assignment_id FROM exam_attempts)""",
            (exam_id,),
        )
        conn.commit()
    finally:
        conn.close()

    payload = _build_payload(db_path, exam_id)
    conn = write_connection(db_path)
    try:
        conn.execute("BEGIN IMMEDIATE")
        existing = conn.execute(
            "SELECT * FROM exam_report_snapshots WHERE exam_event_id=? ORDER BY revision DESC LIMIT 1",
            (exam_id,),
        ).fetchone()
        if existing:
            conn.rollback()
            return {"revision": existing["revision"], "payload": json.loads(existing["payload_json"])}
        now = time_policy.utc_now_ms()
        conn.execute(
            """INSERT INTO exam_report_snapshots
               (id, exam_event_id, revision, schema_version, report_algorithm_version,
                payload_json, checksum, previous_revision, created_by, created_at_ms)
               VALUES (?, ?, 1, ?, ?, ?, ?, NULL, ?, ?)""",
            (gen_id("report"), exam_id, REPORT_SCHEMA_VERSION, REPORT_ALGORITHM_VERSION,
             json.dumps(payload, ensure_ascii=False), _checksum(payload), actor, now),
        )
        conn.execute(
            "UPDATE exam_events SET status='closed', finalized_at_ms=?, finalized_by=? WHERE id=?",
            (now, actor, exam_id),
        )
        write_audit(conn, actor=actor, unit_code=unit_code, action="finalize_exam",
                    entity_type="exam_event", entity_id=exam_id)
        conn.commit()
        return {"revision": 1, "payload": payload}
    finally:
        conn.close()


def get_report_snapshot(db_path, exam_id, revision=None):
    conn = read_connection(db_path)
    try:
        if revision is None:
            row = conn.execute(
                "SELECT revision, payload_json FROM exam_report_snapshots "
                "WHERE exam_event_id=? ORDER BY revision DESC LIMIT 1",
                (exam_id,),
            ).fetchone()
        else:
            row = conn.execute(
                "SELECT revision, payload_json FROM exam_report_snapshots "
                "WHERE exam_event_id=? AND revision=?",
                (exam_id, revision),
            ).fetchone()
        if not row:
            return None
        return {"revision": row["revision"], "payload": json.loads(row["payload_json"])}
    finally:
        conn.close()


def _excel_text(value):
    if value is None:
        return ""
    value = str(value)
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def export_report_excel(payload, output_path):
    """Xuất report snapshot; escape công thức Excel từ dữ liệu người dùng/AI."""
    from openpyxl import Workbook

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("Summary")
    summary.append(["Chỉ tiêu", "Giá trị"])
    for key, value in payload.get("summary", {}).items():
        summary.append([_excel_text(key), value])
    individual = workbook.create_sheet("Individual")
    individual.append(["Username", "Họ tên", "Điểm", "Tỷ lệ", "Đạt"])
    for row in payload.get("individual", []):
        assignment, result = row.get("assignment", {}), row.get("result") or {}
        individual.append([
            _excel_text(assignment.get("username")), _excel_text(assignment.get("display_name")),
            result.get("raw_score", ""), result.get("percent", ""), result.get("passed", ""),
        ])
    for name in ("Topics", "Questions", "Retake"):
        workbook.create_sheet(name)
    workbook.save(output_path)
    return output_path
