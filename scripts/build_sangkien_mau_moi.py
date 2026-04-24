#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_default_font(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)


def set_run_font(run, *, size=13, bold=False, italic=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    align=None,
    bold=False,
    italic=False,
    size=13,
    spacing_after=4,
    spacing_before=0,
    first_line_cm=0,
):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(spacing_after)
    fmt.space_before = Pt(spacing_before)
    if first_line_cm:
        fmt.first_line_indent = Cm(first_line_cm)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return paragraph


def set_cell_text(cell, text: str, *, bold=False, italic=False, align=None, size=12) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    for idx, part in enumerate(str(text).split("\n")):
        if idx > 0:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=bold, italic=italic)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "000000")
        borders.append(elem)
    tbl_pr.append(borders)


def remove_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        borders.append(elem)
    tc_pr.append(borders)


def build_doc(data: dict, template: str | None = None) -> Document:
    doc = Document(template) if template else Document()
    clear_document(doc)
    set_default_font(doc)

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    add_paragraph(
        doc,
        data["title_line_1"],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=15,
        spacing_after=2,
    )
    add_paragraph(
        doc,
        data["title_line_2"],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        italic=False,
        size=13,
        spacing_after=10,
    )
    add_paragraph(doc, f"Kính gửi: {data['recipient']}", spacing_after=6)
    add_paragraph(doc, "Chúng tôi ghi tên dưới đây:", spacing_after=4)

    authors = data["authors"]
    author_table = doc.add_table(rows=len(authors) + 1, cols=8)
    author_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    author_table.style = "Table Grid"
    set_table_borders(author_table)
    headers = [
        "TT",
        "Họ tên tác giả",
        "Nam/Nữ",
        "Trình độ chuyên môn",
        "Chức vụ, đơn vị công tác",
        "Chủ trì / SK",
        "Tỷ lệ đóng góp (%)",
        "Ký tên",
    ]
    for index, header in enumerate(headers):
        set_cell_text(author_table.rows[0].cells[index], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_idx, author in enumerate(authors, start=1):
        values = [
            author.get("index", str(row_idx)),
            author.get("name", ""),
            author.get("gender", ""),
            author.get("education", ""),
            author.get("role_unit", ""),
            author.get("lead", ""),
            author.get("contribution", ""),
            author.get("signature", ""),
        ]
        for col_idx, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in (0, 2, 5, 6, 7) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(author_table.rows[row_idx].cells[col_idx], value, align=align)

    add_paragraph(doc, data["contribution_note"], italic=True, size=12, spacing_before=4, spacing_after=8)
    add_paragraph(doc, f"Điện thoại: {data['phone']} - Email: {data['email']}", spacing_after=4)
    add_paragraph(doc, f"Địa chỉ bưu điện: {data['address']}", spacing_after=4)
    add_paragraph(doc, data["legal_basis"], spacing_after=4)
    add_paragraph(doc, data["proposal_line"], spacing_after=4)
    add_paragraph(doc, data["title_note"], italic=True, size=12, spacing_after=6)
    add_paragraph(doc, data["start_date_line"], spacing_after=4)
    add_paragraph(doc, data["location_line"], spacing_after=10)
    add_paragraph(
        doc,
        data["main_heading"],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=14,
        spacing_after=8,
    )

    for section_data in data["sections"]:
        add_paragraph(doc, section_data["heading"], bold=True, spacing_after=4)
        for paragraph in section_data["paragraphs"]:
            add_paragraph(doc, paragraph, first_line_cm=1, spacing_after=4)
        add_paragraph(doc, "", spacing_after=2)

    add_paragraph(doc, data["results_heading"], bold=True, spacing_after=6)
    results = data["results_rows"]
    results_table = doc.add_table(rows=len(results) + 1, cols=3)
    results_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    results_table.style = "Table Grid"
    set_table_borders(results_table)
    result_headers = [
        "",
        "Mô tả đối tượng / trước khi áp dụng sáng kiến",
        "Mô tả đối tượng / sau khi áp dụng sáng kiến / (hiệu quả kinh tế, lợi ích đã đạt được về các mặt: kinh tế, kỹ thuật, xã hội, môi trường..)",
    ]
    for index, header in enumerate(result_headers):
        set_cell_text(results_table.rows[0].cells[index], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_idx, row in enumerate(results, start=1):
        set_cell_text(results_table.rows[row_idx].cells[0], row["label"], align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(results_table.rows[row_idx].cells[1], row["before"])
        set_cell_text(results_table.rows[row_idx].cells[2], row["after"])

    add_paragraph(doc, "", spacing_after=4)
    add_paragraph(doc, data["self_eval_heading"], bold=True, spacing_after=6)
    eval_rows = data["self_eval_rows"]
    eval_table = doc.add_table(rows=len(eval_rows) + 1, cols=4)
    eval_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    eval_table.style = "Table Grid"
    set_table_borders(eval_table)
    eval_headers = [
        "TT",
        "Tiêu chí",
        "Hướng dẫn cách tự đánh giá",
        "Tác giả tự đánh giá",
    ]
    for index, header in enumerate(eval_headers):
        set_cell_text(eval_table.rows[0].cells[index], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_idx, row in enumerate(eval_rows, start=1):
        set_cell_text(eval_table.rows[row_idx].cells[0], row.get("tt", ""), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(eval_table.rows[row_idx].cells[1], row.get("criterion", ""))
        set_cell_text(eval_table.rows[row_idx].cells[2], row.get("guide", ""))
        set_cell_text(eval_table.rows[row_idx].cells[3], row.get("author_eval", ""))

    add_paragraph(doc, "", spacing_after=6)
    add_paragraph(doc, data["declaration"], spacing_after=10)

    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in sign_table.rows:
        for cell in row.cells:
            remove_cell_borders(cell)
    set_cell_text(sign_table.cell(0, 0), data["sign_left"], bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(sign_table.cell(0, 1), data["sign_right"], bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)

    return doc


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Path to JSON spec")
    parser.add_argument("--output", required=True, help="Output DOCX path")
    parser.add_argument("--template", help="Optional DOCX template/style seed")
    args = parser.parse_args()

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    doc = build_doc(data, template=args.template)
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
