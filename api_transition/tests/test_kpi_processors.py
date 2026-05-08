from pathlib import Path
import sys

import pandas as pd
from openpyxl import Workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from api_transition.processors.common import configure_runtime_roots, reset_runtime_roots
from api_transition.processors.kpi_processors import process_kpi_nvkt_c11_api_output


def _write_c11_kpi_workbook(path):
    wb = Workbook()
    ws = wb.active

    headers = [
        "NVKT_RAW",
        "SM1",
        "SM2",
        "TY_LE_CHU_DONG",
        "SM3",
        "SM4",
        "TY_LE_BAO_HONG",
        "SM5",
        "SM6",
        "TY_LE_CCCO",
        "BSC",
    ]
    for col_index, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_index, value=header)
        ws.cell(row=2, column=col_index, value=header)

    values = [
        "KV-001-Nguyễn Văn A",
        10,
        9,
        90,
        20,
        18,
        91,
        30,
        27,
        92,
        93,
    ]
    for col_index, value in enumerate(values, start=1):
        ws.cell(row=3, column=col_index, value=value)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_process_kpi_nvkt_c11_preserves_new_ccco_columns(tmp_path):
    downloads_root = tmp_path / "downloads"
    processed_root = tmp_path / "Processed"
    input_path = downloads_root / "kpi_nvkt" / "c11-nvktdb report.xlsx"
    dsnv_path = tmp_path / "dsnv.xlsx"

    _write_c11_kpi_workbook(input_path)
    pd.DataFrame(
        [{"Họ tên": "Nguyễn Văn A", "đơn vị": "Tổ VT 1"}]
    ).to_excel(dsnv_path, index=False)

    configure_runtime_roots(downloads_root=downloads_root, processed_root=processed_root)
    try:
        processed_path = process_kpi_nvkt_c11_api_output(
            input_path=input_path,
            dsnv_file=dsnv_path,
            overwrite_processed=True,
        )
    finally:
        reset_runtime_roots()

    result = pd.read_excel(processed_path, sheet_name="c11 kpi nvkt")

    assert list(result.columns) == [
        "STT",
        "đơn vị",
        "NVKT",
        "SM1",
        "SM2",
        "Tỷ lệ sửa chữa phiếu chất lượng chủ động dịch vụ FiberVNN, MyTV đạt yêu cầu",
        "SM3",
        "SM4",
        "Tỷ lệ phiếu sửa chữa báo hỏng dịch vụ BRCĐ đúng quy định không tính hẹn",
        "SM5",
        "SM6",
        "Tỷ lệ phiếu sửa chữa trong ngày tại CCCO",
        "Chỉ tiêu BSC",
    ]
    assert result.loc[0, "SM5"] == 30
    assert result.loc[0, "SM6"] == 27
    assert result.loc[0, "Tỷ lệ phiếu sửa chữa trong ngày tại CCCO"] == 92
    assert result.loc[0, "Chỉ tiêu BSC"] == 93
